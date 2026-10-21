"""Official African sovereign-yield sources with bounded public downloads."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from html import unescape
import json
import re
from urllib.parse import unquote, urljoin

from openpyxl import load_workbook


BOTSWANA_PAGE = 'https://www.bankofbotswana.bw/content/government-securities'
ZAMBIA_API = (
    'https://www.boz.zm/jsonapi/node/bond_auction'
    '?sort=-field_bond_auction_date&page%5Blimit%5D=5'
    '&include=field_auction_results,field_results_file'
)
TANZANIA_PAGE = 'https://www.bot.go.tz/TBonds?lang=en'
TANZANIA_RESULTS = 'https://www.bot.go.tz/TBonds/AuctionSummaries'


def _latest_botswana_workbook(page: str) -> str:
    links = re.findall(r'''href=["']([^"']*Bond-Yields[^"']*\.xlsx(?:\?[^"']*)?)["']''', page, re.I)
    if not links:
        raise ValueError('Bank of Botswana daily bond-yield workbook link was not found')

    def published_date(link: str):
        decoded = unquote(link)
        match = re.search(r'(\d{1,2})[ _-]+([A-Za-z]+)[ _-]+(\d{4})', decoded)
        if not match:
            return date.min
        try:
            return datetime.strptime(' '.join(match.groups()), '%d %B %Y').date()
        except ValueError:
            return date.min

    return urljoin(BOTSWANA_PAGE, max(links, key=published_date))


def botswana(download, download_binary, curve, today, supported_months):
    """Read the latest dated Bank of Botswana daily securities workbook."""
    page = download(BOTSWANA_PAGE)
    url = _latest_botswana_workbook(page)
    # The workbook is small; normal mode keeps random cell reads efficient while
    # scanning its historical rows and maturity headers.
    workbook = load_workbook(BytesIO(download_binary(url)), read_only=False, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        headers = [str(sheet.cell(3, column).value or '').strip() for column in range(1, sheet.max_column + 1)]
        bill_columns = {}
        bond_columns = []
        for column, header in enumerate(headers, start=1):
            bill = re.fullmatch(r'(\d+)\s+months?\s+t[ -]?bill', header, re.I)
            if bill:
                months = int(bill[1])
                if months in supported_months:
                    bill_columns[column] = months
                continue

            maturity = sheet.cell(6, column).value
            if header and isinstance(maturity, (datetime, date)) and column > 1:
                maturity = maturity.date() if isinstance(maturity, datetime) else maturity
                if maturity > today:
                    bond_columns.append((column, maturity))

        rows = []
        for row in range(10, sheet.max_row + 1):
            observed = sheet.cell(row, 1).value
            if not isinstance(observed, (datetime, date)):
                continue
            observed = observed.date() if isinstance(observed, datetime) else observed
            if observed <= today:
                rows.append((observed, row))
        if not rows:
            raise ValueError('Bank of Botswana workbook has no dated observations')
        period, row = max(rows, key=lambda item: item[0])

        values = {}
        for column, months in bill_columns.items():
            raw = sheet.cell(row, column).value
            try:
                value = float(raw)
            except (TypeError, ValueError):
                continue
            if value == value and abs(value) != float('inf'):
                values[months] = value

        # Map an on-the-run bond to a supported point only when its actual
        # remaining term is within 45 days of that existing maturity.
        for column, maturity in bond_columns:
            raw = sheet.cell(row, column).value
            try:
                value = float(raw)
            except (TypeError, ValueError):
                continue
            if value != value or abs(value) == float('inf'):
                continue
            days = (maturity - period).days
            nearest = min(supported_months, key=lambda months: abs(days - months * 365.2425 / 12))
            if abs(days - nearest * 365.2425 / 12) <= 45 and nearest not in values:
                values[nearest] = value

        if len(values) < 2:
            raise ValueError('Bank of Botswana workbook has fewer than two supported yields')
        return {'BWA': curve(
            'BWP', period.isoformat(), 'Bank of Botswana · daily government securities yields', url,
            'Daily 3-, 6- and 12-month Treasury-bill yields, plus on-the-run bond yields only when the issue’s actual remaining maturity is within 45 days of an existing app tenor. Other maturities are omitted.',
            values,
        )}
    finally:
        workbook.close()


def zambia(download, curve, today, supported_months):
    """Use yields reported for tenors in the latest complete Bank of Zambia auction."""
    payload = json.loads(download(ZAMBIA_API))
    included = {item.get('id'): item for item in payload.get('included', [])}
    for auction in payload.get('data', []):
        attributes = auction.get('attributes', {})
        raw_date = attributes.get('field_bond_auction_date')
        if not raw_date:
            continue
        try:
            observed = datetime.fromisoformat(raw_date.replace('Z', '+00:00')).date()
        except ValueError:
            continue
        if observed > today:
            continue

        relationships = auction.get('relationships', {}).get('field_auction_results', {}).get('data', [])
        values = {}
        for relationship in relationships:
            bond = included.get(relationship.get('id'), {})
            match = re.fullmatch(r'paragraph--(\d+)_year_bond', bond.get('type', ''))
            if not match:
                continue
            months = int(match[1]) * 12
            if months not in supported_months:
                continue
            raw_yield = next((value for key, value in bond.get('attributes', {}).items()
                              if key.endswith('yield_rate')), None)
            try:
                value = float(raw_yield)
            except (TypeError, ValueError):
                continue
            if value == value and abs(value) != float('inf'):
                values[months] = value

        if len(values) >= 2:
            auction_id = attributes.get('title', 'latest auction')
            return {'ZMB': curve(
                'ZMW', observed.isoformat(), f'Bank of Zambia · government bond auction {auction_id}', ZAMBIA_API,
                'Primary-market yield rates reported by the Bank of Zambia for bonds tendered in one auction; exact issue tenors are shown. This is an auction cross-section, not a secondary-market constant-maturity curve.',
                values, 'auction',
            )}
    raise ValueError('Bank of Zambia returned no complete dated auction with two supported bond yields')


def tanzania(download, download_form, curve, today, supported_months):
    """Build a monthly auction cross-section from Bank of Tanzania results."""
    html = download(TANZANIA_PAGE)
    auctions = {}
    for row_match in re.finditer(r'<tr\b[^>]*>(.*?)</tr>', html, re.I | re.S):
        row = row_match[1]
        cells = re.findall(r'<td\b[^>]*>(.*?)</td>', row, re.I | re.S)
        button = re.search(
            r'<button\b[^>]*id=["\']showSummaryDetails["\'][^>]*value=["\']([^"\']+)',
            row, re.I | re.S,
        )
        if len(cells) < 3 or not button:
            continue
        title = unescape(re.sub(r'<[^>]+>', ' ', cells[1])).strip()
        raw_date = unescape(re.sub(r'<[^>]+>', ' ', cells[2])).strip()
        tenor = re.search(r'(\d+)\s*[- ]?YEAR(?:S)?\s+TREASURY BOND', title, re.I)
        if not tenor:
            continue
        months = int(tenor[1]) * 12
        if months not in supported_months:
            continue
        try:
            observed = datetime.strptime(raw_date, '%d-%b-%Y').date()
        except ValueError:
            continue
        if observed > today:
            continue
        auction_number, auction_days = button[1].split('_', 1)
        month = observed.strftime('%Y-%m')
        record = dict(date=observed, months=months,
                      auction_number=auction_number, auction_days=auction_days)
        if months not in auctions.setdefault(month, {}) or observed > auctions[month][months]['date']:
            auctions[month][months] = record

    eligible = [(month, points) for month, points in auctions.items() if len(points) >= 2]
    if not eligible:
        raise ValueError('Bank of Tanzania has no month with two supported bond-auction tenors')
    period, selected = max(eligible, key=lambda item: item[0])

    values = {}
    for months, auction in selected.items():
        response = json.loads(download_form(TANZANIA_RESULTS, {
            'au_no': auction['auction_number'], 'au_days': auction['auction_days'],
        }, TANZANIA_PAGE))
        if str(response.get('message', '')).upper() != 'SUCCESS':
            continue
        try:
            detail_date = datetime.strptime(response.get('BondDate', ''), '%B %d, %Y').date()
        except ValueError:
            continue
        if detail_date != auction['date']:
            continue
        entry = next((item for item in response.get('tbondSummary', [])
                      if 'WEIGHTED AVERAGE YIELD-TO-MATURITY' in str(item.get('itemDesc', '')).upper()), None)
        try:
            value = float(entry.get('itemValue')) if entry else None
        except (TypeError, ValueError):
            continue
        if value is not None and value == value and abs(value) != float('inf'):
            values[months] = value

    if len(values) < 2:
        raise ValueError('Bank of Tanzania returned fewer than two dated auction yields')
    return {'TZA': curve(
        'TZS', period, 'Bank of Tanzania · Treasury bond auction results', TANZANIA_PAGE,
        'Latest weighted-average yield-to-maturity from primary bond auctions for each exact tenor in the observation month. The selected tenor auctions may occur on different days, so this is a monthly auction cross-section, not a same-day secondary-market curve.',
        values, 'auction',
    )}
